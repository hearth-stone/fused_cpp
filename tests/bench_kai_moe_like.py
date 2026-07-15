#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""MoE-like KleidiAI GEMM 压测脚本。

模拟 MoE 场景的工作负载：``num_groups`` 组，每组 2 个连续 GEMM，
    GEMM1: [M, K1] @ [K1, N1]   (默认 K1=768,  N1=5120)
    GEMM2: [M, K2] @ [K2, N2]   (默认 K2=5120, N2=384)
其中每组的 M 从 ``[m_min, m_max]`` 随机选取（组间独立）。

对比两种并行策略 ::

    sequential:  1 个 worker 顺序处理所有 GEMM，每个 handler 使用 total_threads 线程。
    grouped(X):  total_threads 按组数 X 拆分，每组 total_threads // X 线程并行；
                 X 个 worker 并行 "领取" GEMM 任务（动态调度）。

输出 ::
    - CSV：每个策略一行，记录总时长、总 FLOPs、GFLOP/s、每组详细耗时。
    - Excel：汇总页 + 详情页（需要 ``openpyxl``）。

用法示例::

    taskset -c 0-79 \\
    env OMP_NUM_THREADS=1 MKL_NUM_THREADS=1 \\
    python tests/bench_kai_moe_like.py \\
        --num-groups 160 --m-min 50 --m-max 150 \\
        --total-threads 80 --group-splits 1 2 4 5 8 10 20 \\
        --warmup 1 --repeat 3 \\
        --output-csv bench_moe.csv \\
        --output-xlsx bench_moe.xlsx

注意：
    - ``sequential`` 策略 handler 内用 total_threads 线程；
    - ``grouped(X)`` 的 X 必须能整除 total_threads；
    - 该脚本不主动 pin CPU，由外部 ``taskset`` 负责核心绑定。
"""

import argparse
import csv
import os
import platform
import random
import sys
import threading
import time
from dataclasses import dataclass, field
from typing import Dict, List, Tuple

import torch


# ── 数据结构 ──


@dataclass
class GEMMSpec:
    """单次 GEMM 规格。"""

    group_id: int
    stage: int  # 0 = GEMM1, 1 = GEMM2
    m: int
    k: int
    n: int

    def flops(self) -> float:
        return 2.0 * self.m * self.k * self.n


@dataclass
class StrategyResult:
    """一次策略运行的聚合结果。"""

    strategy: str
    total_threads: int
    group_splits: int  # 1 表示 sequential
    threads_per_handler: int
    num_gemms: int
    total_flops: float
    wall_time_ms: float
    gflops: float
    per_gemm_ms: List[float] = field(default_factory=list)


# ── 任务生成 ──


def _build_specs(
    num_groups: int,
    m_min: int,
    m_max: int,
    k1: int,
    n1: int,
    k2: int,
    n2: int,
    seed: int,
) -> List[GEMMSpec]:
    """生成 ``num_groups`` 组、每组 2 个 GEMM 的 spec 列表。"""
    rng = random.Random(seed)
    specs: List[GEMMSpec] = []
    for g in range(num_groups):
        m = rng.randint(m_min, m_max)
        specs.append(GEMMSpec(group_id=g, stage=0, m=m, k=k1, n=n1))
        specs.append(GEMMSpec(group_id=g, stage=1, m=m, k=k2, n=n2))
    return specs


# ── Handler 缓存 ──
# 为了避免多次预打包相同的 (K, N) 权重，我们按 (K, N) 缓存 handler。
# 新架构下 handler 本身不持有线程资源，所有策略下可以安全共享；
# 线程池在 _run_sequential / _run_grouped 里按策略独立创建。


class HandlerCache:
    """按 (K, N) 缓存 KAI handler 与其输入权重。"""

    def __init__(self, fused_cpp_mod) -> None:
        self._mod = fused_cpp_mod
        self._cache: Dict[Tuple[int, int], "object"] = {}
        self._weight_ref: Dict[Tuple[int, int], torch.Tensor] = {}

    def _get_weight(self, k: int, n: int, seed: int) -> torch.Tensor:
        key = (k, n)
        if key not in self._weight_ref:
            gen = torch.Generator().manual_seed(seed ^ (k * 131071 + n))
            self._weight_ref[key] = torch.randn(
                k,
                n,
                dtype=torch.float32,
                generator=gen,
            )
        return self._weight_ref[key]

    def get(self, k: int, n: int, seed: int):
        key = (k, n)
        handler = self._cache.get(key)
        if handler is not None:
            return handler
        weight = self._get_weight(k, n, seed)
        packed, pk, pn = self._mod.kai_gemm_prepare(weight)
        handler = self._mod.create_kai_gemm(packed, pk, pn)
        self._cache[key] = handler
        return handler


# ── 输入张量缓存 ──


def _build_inputs(specs: List[GEMMSpec], seed: int) -> Dict[Tuple[int, int], torch.Tensor]:
    """按 (group_id, stage) 预生成 FP32 输入张量，避免压测循环中分配开销。"""
    out: Dict[Tuple[int, int], torch.Tensor] = {}
    for s in specs:
        gen = torch.Generator().manual_seed(
            seed ^ (s.group_id * 1000003 + s.stage * 19 + s.m),
        )
        out[(s.group_id, s.stage)] = torch.randn(
            s.m,
            s.k,
            dtype=torch.float32,
            generator=gen,
        )
    return out


# ── 策略实现 ──


def _run_sequential(
    specs: List[GEMMSpec],
    inputs: Dict[Tuple[int, int], torch.Tensor],
    cache: HandlerCache,
    total_threads: int,
    seed: int,
    out_dtype: torch.dtype,
    fused_cpp_mod,
) -> Tuple[float, List[float]]:
    """单线程调度，使用单个 total_threads 线程池并行计算每个 GEMM。"""
    import fused_cpp as _fcpp  # 延迟导入，避免顶级循环依赖

    per_gemm_ms: List[float] = []
    pool = _fcpp.KAIThreadPool(list(range(total_threads))) if total_threads > 1 else None
    try:
        t_start = time.perf_counter()
        for s in specs:
            handler = cache.get(s.k, s.n, seed)
            x = inputs[(s.group_id, s.stage)]
            t0 = time.perf_counter()
            _ = fused_cpp_mod.kai_gemm(
                handler,
                x,
                output_dtype=out_dtype,
                pool=pool,
            )
            t1 = time.perf_counter()
            per_gemm_ms.append((t1 - t0) * 1e3)
        t_end = time.perf_counter()
    finally:
        if pool is not None:
            pool.close()
    return (t_end - t_start) * 1e3, per_gemm_ms


def _run_grouped(
    specs: List[GEMMSpec],
    inputs: Dict[Tuple[int, int], torch.Tensor],
    cache_factory,
    total_threads: int,
    group_splits: int,
    seed: int,
    out_dtype: torch.dtype,
    fused_cpp_mod,
) -> Tuple[float, List[float]]:
    """把 total_threads 分成 group_splits 组，每组 total_threads // group_splits 线程。

    采用动态调度：共享一个原子索引，每个 worker 循环 "领取" 下一个 spec 来算。
    这样可以容忍各 GEMM 计算量不等导致的负载不均。
    """
    assert total_threads % group_splits == 0, "total_threads 必须能被 group_splits 整除"
    threads_per_handler = total_threads // group_splits

    # 共享任务索引 + 每任务耗时记录
    next_idx = [0]
    idx_lock = threading.Lock()
    per_gemm_ms: List[float] = [0.0] * len(specs)
    errors: List[BaseException] = []
    err_lock = threading.Lock()

    def _worker(wid: int) -> None:
        import fused_cpp as _fcpp  # 延迟导入

        worker_pool = _fcpp.KAIThreadPool(list(range(threads_per_handler))) if threads_per_handler > 1 else None
        try:
            # 每个 worker 独立的 handler cache（避免为了共享而加锁）。
            local_cache = cache_factory(threads_per_handler)
            while True:
                with idx_lock:
                    idx = next_idx[0]
                    if idx >= len(specs):
                        return
                    next_idx[0] = idx + 1
                s = specs[idx]
                handler = local_cache.get(s.k, s.n, seed)
                x = inputs[(s.group_id, s.stage)]
                t0 = time.perf_counter()
                _ = fused_cpp_mod.kai_gemm(
                    handler,
                    x,
                    output_dtype=out_dtype,
                    pool=worker_pool,
                )
                t1 = time.perf_counter()
                per_gemm_ms[idx] = (t1 - t0) * 1e3
        except BaseException as exc:  # pylint: disable=broad-except
            with err_lock:
                errors.append(exc)
        finally:
            if worker_pool is not None:
                worker_pool.close()

    workers = [threading.Thread(target=_worker, args=(i,), name=f"kai-worker-{i}") for i in range(group_splits)]

    t_start = time.perf_counter()
    for w in workers:
        w.start()
    for w in workers:
        w.join()
    t_end = time.perf_counter()

    if errors:
        raise errors[0]

    return (t_end - t_start) * 1e3, per_gemm_ms


# ── 单次策略压测 ──


def _bench_one(
    strategy: str,
    group_splits: int,
    specs: List[GEMMSpec],
    inputs: Dict[Tuple[int, int], torch.Tensor],
    total_threads: int,
    warmup: int,
    repeat: int,
    seed: int,
    out_dtype: torch.dtype,
    fused_cpp_mod,
) -> StrategyResult:
    """执行一次策略压测（含 warmup + repeat），返回最佳 wall time。"""
    total_flops = sum(s.flops() for s in specs)
    threads_per_handler = total_threads if strategy == "sequential" else total_threads // group_splits

    if strategy == "sequential":
        cache = HandlerCache(fused_cpp_mod)

        def _one_run() -> Tuple[float, List[float]]:
            return _run_sequential(
                specs,
                inputs,
                cache,
                total_threads,
                seed,
                out_dtype,
                fused_cpp_mod,
            )
    elif strategy == "grouped":

        def _cache_factory(_threads: int) -> HandlerCache:
            return HandlerCache(fused_cpp_mod)

        def _one_run() -> Tuple[float, List[float]]:
            return _run_grouped(
                specs,
                inputs,
                _cache_factory,
                total_threads,
                group_splits,
                seed,
                out_dtype,
                fused_cpp_mod,
            )
    else:
        raise ValueError(f"未知策略 {strategy!r}")

    # warmup
    for _ in range(warmup):
        _one_run()

    best_wall_ms = float("inf")
    best_per_gemm: List[float] = []
    for _ in range(repeat):
        wall_ms, per_gemm = _one_run()
        if wall_ms < best_wall_ms:
            best_wall_ms = wall_ms
            best_per_gemm = per_gemm

    gflops = total_flops / (best_wall_ms / 1e3) / 1e9 if best_wall_ms > 0 else 0.0
    return StrategyResult(
        strategy=strategy,
        total_threads=total_threads,
        group_splits=group_splits if strategy == "grouped" else 1,
        threads_per_handler=threads_per_handler,
        num_gemms=len(specs),
        total_flops=total_flops,
        wall_time_ms=best_wall_ms,
        gflops=gflops,
        per_gemm_ms=best_per_gemm,
    )


# ── 结果写出 ──


def _write_csv(
    results: List[StrategyResult],
    specs: List[GEMMSpec],
    path: str,
) -> None:
    with open(path, "w", newline="", encoding="utf-8") as fp:
        writer = csv.writer(fp)
        writer.writerow(
            [
                "strategy",
                "total_threads",
                "group_splits",
                "threads_per_handler",
                "num_gemms",
                "total_flops",
                "wall_time_ms",
                "gflops",
            ]
        )
        for r in results:
            writer.writerow(
                [
                    r.strategy,
                    r.total_threads,
                    r.group_splits,
                    r.threads_per_handler,
                    r.num_gemms,
                    f"{r.total_flops:.3e}",
                    f"{r.wall_time_ms:.3f}",
                    f"{r.gflops:.2f}",
                ]
            )
        writer.writerow([])
        writer.writerow(["-- per-GEMM latency (ms), best repeat --"])
        header = ["gemm_idx", "group_id", "stage", "M", "K", "N"] + [
            f"{r.strategy}/x={r.group_splits}" for r in results
        ]
        writer.writerow(header)
        for idx, s in enumerate(specs):
            row = [idx, s.group_id, s.stage, s.m, s.k, s.n]
            for r in results:
                row.append(f"{r.per_gemm_ms[idx]:.3f}" if idx < len(r.per_gemm_ms) else "")
            writer.writerow(row)


def _write_xlsx(
    results: List[StrategyResult],
    specs: List[GEMMSpec],
    path: str,
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        print("[WARN] openpyxl 未安装，跳过 Excel 输出", file=sys.stderr)
        return

    wb = Workbook()
    # Summary sheet
    ws = wb.active
    ws.title = "Summary"
    headers = [
        "strategy",
        "total_threads",
        "group_splits",
        "threads_per_handler",
        "num_gemms",
        "total_flops",
        "wall_time_ms",
        "gflops",
        "speedup_vs_seq",
    ]
    for col_i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col_i, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FFDDDDDD")
        c.alignment = Alignment(horizontal="center")

    seq = next((r for r in results if r.strategy == "sequential"), None)
    base_ms = seq.wall_time_ms if seq else None
    for row_i, r in enumerate(results, start=2):
        ws.cell(row=row_i, column=1, value=r.strategy)
        ws.cell(row=row_i, column=2, value=r.total_threads)
        ws.cell(row=row_i, column=3, value=r.group_splits)
        ws.cell(row=row_i, column=4, value=r.threads_per_handler)
        ws.cell(row=row_i, column=5, value=r.num_gemms)
        ws.cell(row=row_i, column=6, value=r.total_flops)
        ws.cell(row=row_i, column=7, value=round(r.wall_time_ms, 3))
        ws.cell(row=row_i, column=8, value=round(r.gflops, 2))
        speedup = base_ms / r.wall_time_ms if base_ms and r.wall_time_ms > 0 else None
        if speedup is not None:
            ws.cell(row=row_i, column=9, value=round(speedup, 3))

    for col_letter in ("A", "B", "C", "D", "E", "F", "G", "H", "I"):
        ws.column_dimensions[col_letter].width = 18

    # Detail sheet
    ws2 = wb.create_sheet(title="PerGEMM")
    detail_headers = ["gemm_idx", "group_id", "stage", "M", "K", "N"] + [
        f"{r.strategy}/x={r.group_splits}" for r in results
    ]
    for col_i, h in enumerate(detail_headers, start=1):
        c = ws2.cell(row=1, column=col_i, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FFDDDDDD")
        c.alignment = Alignment(horizontal="center")

    for idx, s in enumerate(specs):
        ws2.cell(row=idx + 2, column=1, value=idx)
        ws2.cell(row=idx + 2, column=2, value=s.group_id)
        ws2.cell(row=idx + 2, column=3, value=s.stage)
        ws2.cell(row=idx + 2, column=4, value=s.m)
        ws2.cell(row=idx + 2, column=5, value=s.k)
        ws2.cell(row=idx + 2, column=6, value=s.n)
        for col_i, r in enumerate(results, start=7):
            if idx < len(r.per_gemm_ms):
                ws2.cell(
                    row=idx + 2,
                    column=col_i,
                    value=round(r.per_gemm_ms[idx], 3),
                )

    for col_letter in ("A", "B", "C", "D", "E", "F"):
        ws2.column_dimensions[col_letter].width = 12

    wb.save(path)


def _print_brief_summary(results: List[StrategyResult]) -> None:
    """终端仅打印一个非常简洁的 summary，详细数据请查看 CSV/Excel。"""
    print()
    print("=" * 96)
    print(f"{'strategy':<14}{'threads':>10}{'x':>6}{'t/handler':>12}{'wall_ms':>12}{'GFLOP/s':>12}{'speedup':>12}")
    print("=" * 96)
    seq = next((r for r in results if r.strategy == "sequential"), None)
    base_ms = seq.wall_time_ms if seq else None
    for r in results:
        speedup = f"{base_ms / r.wall_time_ms:.2f}x" if base_ms and r.wall_time_ms > 0 else "-"
        print(
            f"{r.strategy:<14}{r.total_threads:>10}{r.group_splits:>6}"
            f"{r.threads_per_handler:>12}{r.wall_time_ms:>12.2f}"
            f"{r.gflops:>12.2f}{speedup:>12}"
        )
    print("=" * 96)


# ── 主入口 ──


def build_argparser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--num-groups", type=int, default=160, help="组数（每组 2 个 GEMM），默认 160")
    parser.add_argument("--m-min", type=int, default=50, help="M 随机范围下界，默认 50")
    parser.add_argument("--m-max", type=int, default=150, help="M 随机范围上界，默认 150")
    parser.add_argument("--k1", type=int, default=768, help="GEMM1 的 K，默认 768")
    parser.add_argument("--n1", type=int, default=5120, help="GEMM1 的 N，默认 5120")
    parser.add_argument("--k2", type=int, default=5120, help="GEMM2 的 K，默认 5120")
    parser.add_argument("--n2", type=int, default=384, help="GEMM2 的 N，默认 384")
    parser.add_argument(
        "--total-threads", type=int, default=80, help="总可用线程数（应与 taskset 绑定的核心数一致），默认 80"
    )
    parser.add_argument(
        "--group-splits",
        nargs="+",
        type=int,
        default=[1, 2, 4, 5, 8, 10, 20],
        help="grouped 策略的分组数列表，必须整除 total_threads；值为 1 时会 skip（与 sequential 等价）",
    )
    parser.add_argument(
        "--include-sequential", action="store_true", default=True, help="是否包含 sequential 策略（默认启用）"
    )
    parser.add_argument("--no-sequential", dest="include_sequential", action="store_false", help="关闭 sequential 策略")
    parser.add_argument("--dtype", default="bf16", choices=["fp32", "bf16"], help="输出 dtype，默认 bf16")
    parser.add_argument("--warmup", type=int, default=1, help="warmup 轮数，默认 1")
    parser.add_argument("--repeat", type=int, default=3, help="正式测量轮数，取最短 wall time，默认 3")
    parser.add_argument("--seed", type=int, default=0, help="随机种子，默认 0")
    parser.add_argument("--output-csv", default="bench_kai_moe_like.csv", help="CSV 输出路径")
    parser.add_argument("--output-xlsx", default="bench_kai_moe_like.xlsx", help="Excel 输出路径（需要 openpyxl）")
    return parser


def main() -> int:
    args = build_argparser().parse_args()

    torch.manual_seed(args.seed)
    random.seed(args.seed)

    # 避免 PyTorch intra-op 与 KAI 线程池嵌套导致过度订阅。
    torch.set_num_threads(1)

    try:
        import fused_cpp  # type: ignore[import-not-found]
    except ImportError as exc:
        print(f"[ERROR] 无法导入 fused_cpp：{exc}", file=sys.stderr)
        return 1

    kai_ok = getattr(fused_cpp, "_supports_kai", False)
    is_aarch64 = platform.machine() in ("aarch64", "arm64")
    if not (kai_ok and is_aarch64):
        print("[ERROR] KleidiAI 后端不可用（需要 AArch64 且已启用 KleidiAI 构建）", file=sys.stderr)
        return 2

    # 校验 group_splits
    for x in args.group_splits:
        if x < 1 or args.total_threads % x != 0:
            print(f"[ERROR] group_splits 中的 {x} 无法整除 total_threads={args.total_threads}", file=sys.stderr)
            return 3

    dtype_map = {"fp32": torch.float32, "bf16": torch.bfloat16}
    out_dtype = dtype_map[args.dtype]

    specs = _build_specs(
        args.num_groups,
        args.m_min,
        args.m_max,
        args.k1,
        args.n1,
        args.k2,
        args.n2,
        args.seed,
    )
    inputs = _build_inputs(specs, args.seed)

    env_info = {
        "pid": os.getpid(),
        "cpu_count": os.cpu_count(),
        "machine": platform.machine(),
        "torch_num_threads": torch.get_num_threads(),
        "OMP_NUM_THREADS": os.environ.get("OMP_NUM_THREADS", "<unset>"),
    }
    print(f"[INFO] env: {env_info}")
    print(f"[INFO] 共 {len(specs)} 个 GEMM，total_threads={args.total_threads}，dtype={args.dtype}")

    results: List[StrategyResult] = []

    if args.include_sequential:
        print("[INFO] 运行 sequential 策略 ...")
        r = _bench_one(
            "sequential",
            1,
            specs,
            inputs,
            args.total_threads,
            args.warmup,
            args.repeat,
            args.seed,
            out_dtype,
            fused_cpp,
        )
        print(f"       wall={r.wall_time_ms:.2f} ms, {r.gflops:.2f} GFLOP/s")
        results.append(r)

    for x in args.group_splits:
        if x == 1:
            # 与 sequential 等价，无意义，跳过。
            continue
        print(f"[INFO] 运行 grouped(x={x}) 策略 ({args.total_threads // x} threads/handler) ...")
        r = _bench_one(
            "grouped",
            x,
            specs,
            inputs,
            args.total_threads,
            args.warmup,
            args.repeat,
            args.seed,
            out_dtype,
            fused_cpp,
        )
        print(f"       wall={r.wall_time_ms:.2f} ms, {r.gflops:.2f} GFLOP/s")
        results.append(r)

    _write_csv(results, specs, args.output_csv)
    _write_xlsx(results, specs, args.output_xlsx)

    _print_brief_summary(results)
    print(f"\n[OK] CSV  -> {args.output_csv}")
    print(f"[OK] XLSX -> {args.output_xlsx}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
